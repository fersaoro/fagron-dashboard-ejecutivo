"""
Fagron - Dashboard Ejecutivo de Ventas
Script de generacion de data.json a partir de REPORTE_DE_VENTAS.xlsx

Version 3: soporta que el "Resumen ejecutivo" y la tabla de "Cumplimiento"
respondan dinamicamente a los filtros (fecha, zona, linea, asesor, cliente,
familia) en vez de mostrar solo el mes de corte fijo. Para eso se exporta
la cuota de CADA representante por CADA mes del anio (no solo el mes en
curso), y las ventas semanales por asesor de TODOS los meses (no solo el
actual). El calculo de "realizado" para cualquier rango de fechas se hace
siempre en el navegador a partir de "Data 2025-2026" (fuente unica de
verdad para ventas reales); la "cuota" siempre viene de "CUOTAS 2026".

USO SEMANAL:
    python3 build_data.py /ruta/a/REPORTE_DE_VENTAS.xlsx /ruta/salida/data.json
"""
import sys
import json
import datetime
import unicodedata
import openpyxl
from collections import OrderedDict, defaultdict

def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if v in ('', '#N/A', 'N/A'):
            return None
        return v
    return v

def to_num(v):
    if isinstance(v, (int, float)):
        return v
    return 0

def norm_name(v):
    if v is None:
        return None
    v = str(v).strip().upper()
    v = unicodedata.normalize('NFKD', v).encode('ascii', 'ignore').decode('ascii')
    return v

def canon_case(v):
    """Normaliza mayusculas/minusculas para evitar que el mismo asesor o la
    misma familia queden divididos en dos por un error de digitacion en el
    Excel (ej. 'Marly Paola Bolaño Suarez' vs 'MARLY PAOLA BOLAÑO SUAREZ').
    Conserva tildes/enies, solo unifica mayus/minus."""
    return v.strip().upper() if v else v


def parse_cuotas(ws):
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    row_meses = rows[1]

    meses = []
    col = 3
    while len(meses) < 12:
        meses.append((clean(row_meses[col]), col))
        col += 4
    total_start = col

    def bloque(row, start):
        return {
            "cuota": to_num(row[start]),
            "realizado": to_num(row[start + 1]),
            "cumplimiento": (row[start + 2] if isinstance(row[start + 2], (int, float)) else None),
            "diferencia": (row[start + 3] if isinstance(row[start + 3], (int, float)) else None),
        }

    def cuota_por_mes_de_fila(row):
        return [to_num(row[start]) for (_, start) in meses]

    idx_total_general = None
    idx_facturado_ly = None
    data_rows_end = None
    for i, row in enumerate(rows):
        label = clean(row[2])
        if label and label.strip().upper() == 'TOTAL GENERAL 2026':
            idx_total_general = i
            if data_rows_end is None:
                data_rows_end = i
        if label and label.strip().upper().startswith('FACTURADO'):
            idx_facturado_ly = i

    tg_row = rows[idx_total_general]
    current_month_idx = 0
    for m, (nombre, start) in enumerate(meses):
        b = bloque(tg_row, start)
        if b["cuota"]:
            current_month_idx = m
    mes_nombre, mes_col = meses[current_month_idx]

    def kpis_de_fila(row):
        return bloque(row, mes_col), bloque(row, total_start)

    mes_kpi, ytd_kpi = kpis_de_fila(tg_row)
    if idx_facturado_ly is not None:
        mes_ly, ytd_ly = kpis_de_fila(rows[idx_facturado_ly])
    else:
        mes_ly = ytd_ly = {"cuota": 0, "realizado": 0, "cumplimiento": None, "diferencia": None}

    kpis = {
        "mes_nombre": mes_nombre,
        "anio": 2026,
        "mes": {
            "facturado": mes_kpi["realizado"], "cuota": mes_kpi["cuota"],
            "diferencia": mes_kpi["diferencia"], "cumplimiento": mes_kpi["cumplimiento"],
        },
        "ytd": {
            "facturado": ytd_kpi["realizado"], "cuota": ytd_kpi["cuota"],
            "diferencia": ytd_kpi["diferencia"], "cumplimiento": ytd_kpi["cumplimiento"],
        },
        "mes_anterior_anio": {
            "facturado": mes_ly["realizado"],
            "crecimiento_pct": mes_ly["cumplimiento"],
            "crecimiento_abs": mes_ly["diferencia"],
        },
        "ytd_anterior_anio": {
            "facturado": ytd_ly["realizado"],
            "crecimiento_pct": ytd_ly["cumplimiento"],
            "crecimiento_abs": ytd_ly["diferencia"],
        },
        "cuota_por_mes_total": cuota_por_mes_de_fila(tg_row),
        "meses_nombres": [m[0] for m in meses],
    }

    real_general = []
    current_gerente = None
    ZONA_OVERRIDE = {
        norm_name("VENTAS SAC"): "SAC",
        norm_name("VENTAS EMPLEADOS"): "SAC",
    }
    for i in range(4, data_rows_end):
        row = rows[i]
        g, z, r = clean(row[0]), clean(row[1]), clean(row[2])
        if g is None and z is None and r is None:
            if real_general and real_general[-1]["tipo"] != "espacio":
                real_general.append({"tipo": "espacio"})
            current_gerente = None
            continue
        if r is None:
            continue
        if g:
            current_gerente = g
        z = ZONA_OVERRIDE.get(norm_name(r), z)
        r_up = r.strip().upper()
        tipo = "total_zona" if r_up.startswith("TOTAL") else "representante"
        real_general.append({
            "tipo": tipo,
            "gerente": current_gerente,
            "zona": z,
            "representante": r,
            "cuota_por_mes": cuota_por_mes_de_fila(row),
            "asesor_idx": None,
        })
    while real_general and real_general[-1]["tipo"] == "espacio":
        real_general.pop()

    real_general.append({
        "tipo": "total_general",
        "gerente": None, "zona": None, "representante": "TOTAL GENERAL",
        "cuota_por_mes": cuota_por_mes_de_fila(tg_row),
        "asesor_idx": None,
    })

    return kpis, real_general, current_month_idx, mes_nombre


def build(xlsx_path, out_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    kpis, real_general, current_month_idx, mes_nombre = parse_cuotas(wb['CUOTAS 2026'])

    ws = wb['Data 2025-2026']

    zona_idx, linea_idx, asesor_idx, familia_idx, cliente_idx = (OrderedDict() for _ in range(5))
    def idx_of(d, key):
        if key not in d:
            d[key] = len(d)
        return d[key]

    date_set = set()
    agg = {}
    semanas_por_mes_asesor = defaultdict(lambda: defaultdict(lambda: [0.0]*5))
    total_rows = 0
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        total_rows += 1
        fecha = row[3]
        if not isinstance(fecha, datetime.datetime):
            skipped += 1
            continue
        zona = clean(row[25]) or 'SIN ZONA'
        linea = clean(row[17]) or 'SIN LINEA'
        asesor = canon_case(clean(row[20]) or 'SIN ASESOR')
        familia = canon_case(clean(row[13]) or 'SIN FAMILIA')
        cliente = clean(row[7]) or 'SIN CLIENTE'
        valores = to_num(row[16])
        cantidad = to_num(row[15])
        semana = row[4]

        d = fecha.date()
        date_set.add(d)
        zi, li, ai = idx_of(zona_idx, zona), idx_of(linea_idx, linea), idx_of(asesor_idx, asesor)
        fi, ci = idx_of(familia_idx, familia), idx_of(cliente_idx, cliente)
        key = (d.toordinal(), zi, li, ai, fi, ci)
        if key in agg:
            agg[key][0] += valores
            agg[key][1] += cantidad
        else:
            agg[key] = [valores, cantidad]

        if isinstance(semana, (int, float)):
            wk = int(semana)
            if 1 <= wk <= 5:
                ym = "%04d-%02d" % (d.year, d.month)
                semanas_por_mes_asesor[ym][norm_name(asesor)][wk - 1] += valores

    dates_sorted = sorted(date_set)
    date_ord_to_idx = {d.toordinal(): i for i, d in enumerate(dates_sorted)}

    facts = []
    for (dord, zi, li, ai, fi, ci), (val, cant) in agg.items():
        facts.append([date_ord_to_idx[dord], zi, li, ai, fi, ci, round(val, 2), cant])

    asesor_idx_by_norm = {norm_name(k): v for k, v in asesor_idx.items()}
    for item in real_general:
        if item["tipo"] in ("representante", "total_zona", "total_general"):
            ai = asesor_idx_by_norm.get(norm_name(item["representante"]))
            item["asesor_idx"] = ai

    for item in real_general:
        if item["tipo"] == "representante":
            nn = norm_name(item["representante"])
            semanas_mes = {}
            for ym, por_asesor in semanas_por_mes_asesor.items():
                if nn in por_asesor and any(por_asesor[nn]):
                    semanas_mes[ym] = [round(x, 2) for x in por_asesor[nn]]
            item["semanas_por_mes"] = semanas_mes

    def sumar_semanas(items):
        out = {}
        for it in items:
            for ym, wk in (it.get("semanas_por_mes") or {}).items():
                if ym not in out:
                    out[ym] = [0.0]*5
                for k in range(5):
                    out[ym][k] += wk[k]
        return {ym: [round(x, 2) for x in wk] for ym, wk in out.items()}

    reps = [it for it in real_general if it["tipo"] == "representante"]
    for item in real_general:
        if item["tipo"] == "total_zona" and item.get("zona"):
            miembros = [r for r in reps if r.get("zona") == item["zona"]]
            item["semanas_por_mes"] = sumar_semanas(miembros)
        elif item["tipo"] == "total_general":
            item["semanas_por_mes"] = sumar_semanas(reps)

    data = {
        "generado": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "kpis": kpis,
        "real_general": real_general,
        "dims": {
            "fechas": [d.strftime("%Y-%m-%d") for d in dates_sorted],
            "zonas": list(zona_idx.keys()),
            "lineas": list(linea_idx.keys()),
            "asesores": list(asesor_idx.keys()),
            "familias": list(familia_idx.keys()),
            "clientes": list(cliente_idx.keys()),
        },
        "facts": facts,
        "stats": {
            "filas_leidas": total_rows,
            "filas_sin_fecha": skipped,
            "filas_agregadas": len(facts),
        }
    }

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))

    print("OK ->", out_path)
    print("mes en curso detectado:", mes_nombre, kpis["anio"])
    print("filas leidas:", total_rows, "| descartadas sin fecha:", skipped, "| filas agregadas:", len(facts))
    print("zonas:", len(zona_idx), "lineas:", len(linea_idx), "asesores:", len(asesor_idx),
          "familias:", len(familia_idx), "clientes:", len(cliente_idx), "fechas:", len(dates_sorted))
    print("KPIs mes:", kpis["mes"])
    print("KPIs ytd:", kpis["ytd"])
    print("real_general filas:", len(real_general))
    sin_cruzar = [r["representante"] for r in reps if r["asesor_idx"] is None]
    if sin_cruzar:
        print("AVISO: representantes sin match en ASESOR de Data 2025-2026:", sin_cruzar)

if __name__ == "__main__":
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else "REPORTE_DE_VENTAS.xlsx"
    out_path = sys.argv[2] if len(sys.argv) > 2 else "data.json"
    build(xlsx_path, out_path)
